"""NADT 2.0 项目模型 / 服务测试。"""

import os

import openpyxl

from nadt2_lib import secrets
from nadt2_lib.model import Credential, Project, Template
from nadt2_lib import services


def _project() -> Project:
    p = Project(name="测试方案")
    p.templates["sw.cfg"] = Template("sw.cfg", "sysname {{hostname}}\n"
                                                "vlan {{vlan}}\n"
                                                "password {{svcd_password}}\n",
                                     "SW")
    p.settings.default_template = "sw.cfg"
    p.credentials["svcd_password"] = Credential("svcd_password", "TopSecret!", "管理密码")
    p.devices = [
        {"hostname": "SW-1", "type": "SW", "esn": "ESN-1", "mgmt_ip": "10.0.0.1",
         "mgmt_vlan": "10", "access_vlan": "20", "mask": "255.255.255.0",
         "gateway": "10.0.0.254", "note": "", "_extras": {"vlan": "10 20"}},
        {"hostname": "SW-2", "type": "SW", "esn": "ESN-2", "mgmt_ip": "10.0.0.2",
         "mgmt_vlan": "10", "access_vlan": "20", "mask": "255.255.255.0",
         "gateway": "10.0.0.254", "note": "", "_extras": {"vlan": "10 20"}},
    ]
    return p


def test_project_save_load_roundtrip(tmp_path):
    p = _project()
    path = tmp_path / "demo.nadt"
    p.save(str(path))
    loaded = Project.load(str(path))
    assert loaded.name == "测试方案"
    assert len(loaded.devices) == 2
    assert loaded.settings.default_template == "sw.cfg"
    assert loaded.credentials["svcd_password"].value == "TopSecret!"
    assert loaded.templates["sw.cfg"].text.startswith("sysname {{hostname}}")


def test_project_file_is_json_and_no_plain_secret(tmp_path):
    p = _project()
    path = tmp_path / "demo.nadt"
    p.save(str(path))
    text = path.read_text(encoding="utf-8")
    assert '"format": "nadt-project"' in text
    assert "TopSecret!" not in text


def test_secrets_roundtrip():
    token = secrets.protect("P@ssw0rd")
    assert secrets.unprotect(token) == "P@ssw0rd"
    assert token.startswith(("dpapi:", "plain:"))


def test_cell_aliases_and_filename_rules():
    dev = {"hostname": "SW-1", "esn": "ESN-1", "mgmt_ip": "10.0.0.1",
           "gateway": "10.0.0.254", "access_vlan": "20", "mgmt_vlan": "10",
           "_extras": {}}
    assert Project.cell_get(dev, "ip") == "10.0.0.1"
    Project.cell_set(dev, "ip", "10.0.0.9")
    assert dev["mgmt_ip"] == "10.0.0.9"
    p = _project()
    assert p.filename_base(dev, "esn") == "ESN-1"
    assert p.filename_base(dev, "hostname") == "SW-1"


def test_validator_detects_duplicates_and_missing_vars():
    p = _project()
    p.devices[1]["hostname"] = "SW-1"
    p.templates["sw.cfg"].text = "sysname {{hostname}}\n{{not_defined}}\n"
    issues = services.Validator(p).validate()
    codes = {i.code for i in issues}
    assert "host-dup" in codes
    assert "render" in codes


def test_generator_plan_diff_and_write(tmp_path):
    p = _project()
    p.settings.output_dir = str(tmp_path / "out")
    p.settings.template_subdir = False
    p.settings.generate_lswnet = False
    plan = services.Generator(p).plan()
    assert plan.added == 2
    assert all(f.status == "added" for f in plan.files)
    result = services.Generator(p).write(plan)
    assert result["written"] == 2
    assert os.path.isfile(tmp_path / "out" / "SW-1.cfg")

    plan2 = services.Generator(p).plan()
    assert plan2.unchanged == 2
    p.devices[0]["note"] = "changed"
    p.templates["sw.cfg"].text += "# new line\n"
    plan3 = services.Generator(p).plan()
    assert plan3.changed == 2


def test_diff_masks_secrets():
    p = _project()
    old = "password TopSecret!\n"
    new = "password TopSecret!\nnew line\n"
    text = services.diff_text(old, new, "SW-1.cfg", p)
    assert "TopSecret!" not in text
    assert "new line" in text


def test_batch_operations():
    p = _project()
    services.fill_down(p, [0, 1], ["note"], "1F")
    assert all(d["note"] == "1F" for d in p.devices)
    services.sequence(p, [0, 1], "mgmt_vlan", 10, 10, padding=3, prefix="V")
    assert p.devices[0]["mgmt_vlan"] == "V010"
    assert p.devices[1]["mgmt_vlan"] == "V020"
    n = services.find_replace(p, [0, 1], ["note"], "1F", "2F")
    assert n == 2
    new_idx = services.duplicate_rows(p, [0])
    assert p.devices[new_idx[0]]["hostname"] == "SW-1-copy"
    assert p.devices[new_idx[0]]["esn"] == ""


def test_import_legacy_workbook(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "主页"
    tpl = wb.create_sheet("SW")
    tpl["A1"] = "【SW】"
    tpl["C1"] = "{{sysname}}"
    tpl["D1"] = "{{esn}}"
    tpl["A2"] = "sysname {{sysname}}"
    tpl["C2"] = "SW-1"
    tpl["D2"] = "ESN-1"
    path = tmp_path / "legacy.xlsx"
    wb.save(str(path))
    p = Project.create_default("导入")
    info = services.import_legacy_workbook(p, str(path))
    assert info["templates"] == 1
    assert info["devices"] == 1
    assert p.devices[0]["hostname"] == "SW-1"


def test_import_rows_with_column_mapping():
    p = _project()
    header = ["名字", "地址", "备注"]
    rows = [["SW-9", "10.9.9.9", "9F"], ["SW-10", "10.9.9.10", "10F"]]
    n = services.import_rows(p, header, rows,
                             {0: "hostname", 1: "mgmt_ip", 2: "note"}, "sw.cfg")
    assert n == 2
    assert p.devices[-1]["hostname"] == "SW-10"
    assert p.devices[-1]["mgmt_ip"] == "10.9.9.10"
    assert p.devices[-1]["type"] == "sw"


def test_export_workbook_roundtrip(tmp_path):
    p = _project()
    out = tmp_path / "export.xlsx"
    services.export_workbook(p, str(out))
    assert out.is_file()
    wb = openpyxl.load_workbook(str(out))
    assert "sw" in wb.sheetnames
    ws = wb["sw"]
    assert ws["A1"].value == "【sw】"
    assert ws["C2"].value == "SW-1"


def test_project_autosave(tmp_path):
    p = _project()
    p.path = str(tmp_path / "p.nadt")
    p.write_autosave()
    assert os.path.isfile(p.autosave_path())
    p.clear_autosave()
    assert not os.path.exists(p.autosave_path())
