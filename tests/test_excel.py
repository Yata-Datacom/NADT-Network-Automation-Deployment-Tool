"""Excel（xlsx/xlsm）模板导入测试：parse_xlsm_templates。

用 openpyxl 在 tmp_path 里现场造一个「宏脚本工作流」结构的小工作簿，
不依赖任何真实客户模板文件。
"""

import shutil

import openpyxl
import pytest

import nadt as N

HOST1, HOST2 = "SW-1-1F", "SW-1-2F"


def _build_workbook(tmp_path, name="templates.xlsx"):
    """主页（跳过）/ AP-31U（模板+设备）/ lswnet（映射表，跳过）/ 帮助说明（跳过）。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "主页"
    ws["A1"] = "使用说明：请勿修改本 sheet"

    tpl = wb.create_sheet("AP-31U")
    tpl["A1"] = "【AP-31U】"
    tpl["C1"] = "{{sysname}}"
    tpl["D1"] = "{{mgmt_ip}}"
    tpl["E1"] = "{{mgmt_vlan}}"
    tpl["F1"] = "{{esn}}"
    tpl["A2"] = "sysname {{sysname}}"
    tpl["A3"] = "interface Vlanif{{mgmt_vlan}}"
    tpl["A4"] = "# 分隔行保留"
    tpl["C2"], tpl["D2"], tpl["E2"], tpl["F2"] = HOST1, "192.0.2.11", "11", "0000-1111-2222"
    tpl["C3"], tpl["D3"], tpl["E3"], tpl["F3"] = HOST2, "192.0.2.12", "11", "0000-1111-3333"

    mapping = wb.create_sheet("lswnet")
    mapping["A1"] = "【lswnet】"
    mapping["A2"] = f"esn=0000-1111-2222;cfgfile={HOST1}.cfg;"
    mapping["A3"] = f"esn=0000-1111-3333;cfgfile={HOST2}.cfg;"

    help_ws = wb.create_sheet("帮助说明")
    help_ws["A1"] = "【帮助说明】"
    help_ws["A2"] = "随便写点说明"

    path = tmp_path / name
    wb.save(str(path))
    return path


def test_parse_xlsm_templates_templates_devices_warnings(tmp_path):
    path = _build_workbook(tmp_path)
    templates, devices, warnings = N.parse_xlsm_templates(str(path))
    assert list(templates) == ["AP-31U"]
    assert templates["AP-31U"] == ("sysname {{sysname}}\n"
                                   "interface Vlanif{{mgmt_vlan}}\n"
                                   "# 分隔行保留")
    assert [d["hostname"] for d in devices] == [HOST1, HOST2]
    assert all(d["type"] == "AP-31U" for d in devices)
    assert devices[0]["_extras"] == {"sysname": HOST1, "mgmt_ip": "192.0.2.11",
                                     "mgmt_vlan": "11", "esn": "0000-1111-2222"}
    # 映射表 sheet 必须被跳过（不能被当成模板导入），且给出警告
    assert "lswnet" not in templates
    assert any("lswnet" in w for w in warnings)


def test_parse_xlsm_templates_skips_home_and_help_sheets(tmp_path):
    path = _build_workbook(tmp_path)
    templates, devices, _warnings = N.parse_xlsm_templates(str(path))
    assert "主页" not in templates
    assert "帮助说明" not in templates
    assert len(templates) == 1
    assert len(devices) == 2                       # 主页/帮助说明不产生设备


def test_parse_xlsm_templates_accepts_xlsm_extension(tmp_path):
    """同名内容换成 .xlsm 扩展名同样可读（宏工作簿场景）。"""
    src = _build_workbook(tmp_path)
    dst = tmp_path / "templates.xlsm"
    shutil.copy(str(src), str(dst))
    templates, devices, _warnings = N.parse_xlsm_templates(str(dst))
    assert list(templates) == ["AP-31U"]
    assert len(devices) == 2


def test_excel_devices_flow_into_switch_config(tmp_path, make_dev):
    """Excel 导入 → build_switch_config 端到端：模板占位符全部有值，无残留报错。"""
    path = _build_workbook(tmp_path)
    templates, devices, _warnings = N.parse_xlsm_templates(str(path))
    cfg = N.build_switch_config(templates["AP-31U"], devices[0])
    assert cfg == (f"sysname {HOST1}\n"
                   "interface Vlanif11\n"
                   "# 分隔行保留")
    assert "{{" not in cfg


def test_excel_devices_flow_into_lswnet_and_usb_ini(tmp_path):
    path = _build_workbook(tmp_path)
    _templates, devices, _warnings = N.parse_xlsm_templates(str(path))
    # Excel 里的 ESN 参数列会被提升为标准字段 esn
    assert [d["esn"] for d in devices] == ["0000-1111-2222", "0000-1111-3333"]
    lswnet = N.build_lswnet(devices)
    assert lswnet == (f"esn=0000-1111-2222;cfgfile={HOST1}.cfg;\n"
                      f"esn=0000-1111-3333;cfgfile={HOST2}.cfg;\n")
    ini = N.build_usb_ini(devices)
    assert ini.count("[DEVICE") == 2
    assert ini.endswith(";END DC\r\n")


def test_parse_xlsm_templates_empty_workbook(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "主页"
    path = tmp_path / "empty.xlsx"
    wb.save(str(path))
    assert N.parse_xlsm_templates(str(path)) == ({}, [], [])


def test_parse_xlsm_templates_missing_file_raises(tmp_path):
    with pytest.raises(Exception):
        N.parse_xlsm_templates(str(tmp_path / "nope.xlsx"))
