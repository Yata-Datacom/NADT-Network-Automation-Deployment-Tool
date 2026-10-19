"""方案工作台的数据模型测试：模板变量、标准字段提升、旧工作簿 ESN 合并。"""

import openpyxl

import nadt as N


def test_template_variables_excludes_directives_and_loop_vars():
    text = ("{{hostname}}\n"
            "{{#if access_vlan}}vlan {{access_vlan}}{{#endif}}\n"
            "{{#for n in 1-3}} port {{n}}{{#endfor}}\n"
            "{{ ip }} {{hostname}}")
    assert N.template_variables(text) == ["hostname", "access_vlan", "ip"]


def test_template_variables_ignores_if_literal_value():
    assert N.template_variables("{{#if type=AP}}x{{#endif}}") == ["type"]


def test_build_switch_config_sysname_falls_back_to_hostname():
    cfg = N.build_switch_config("sysname {{sysname}}\n", {"hostname": "SW-1"})
    assert cfg == "sysname SW-1\n"


def test_parse_sheet_rows_promotes_common_standard_fields():
    rows = [
        ["【JR】", None, "{{sysname}}", "{{ip}}", "{{wg}}", "{{access}}", "{{vlanif}}"],
        ["sysname {{sysname}}", None, "SW-1", "10.0.0.11", "10.0.0.254", "20", "10"],
    ]
    _t, devices = N.parse_sheet_rows(rows, "JR", [])
    dev = devices[0]
    assert dev["hostname"] == "SW-1"
    assert dev["mgmt_ip"] == "10.0.0.11"
    assert dev["gateway"] == "10.0.0.254"
    assert dev["access_vlan"] == "20"
    assert dev["mgmt_vlan"] == "10"


def test_parse_xlsm_merges_lswnet_esn_by_hostname(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "主页"
    tpl = wb.create_sheet("JR")
    tpl["A1"] = "【JR】"
    tpl["C1"] = "{{sysname}}"
    tpl["A2"] = "sysname {{sysname}}"
    tpl["C2"] = "SW-1"
    mapping = wb.create_sheet("lswnet")
    mapping["A1"] = "【lswnet】"
    mapping["A2"] = "esn=ESN-001;vrpfile=SW.cc;patchfile=SW.pat;cfgfile=SW-1.cfg;"
    path = tmp_path / "book.xlsx"
    wb.save(str(path))

    templates, devices, _warnings = N.parse_xlsm_templates(str(path))
    assert "JR" in templates
    dev = next(d for d in devices if d["hostname"] == "SW-1")
    assert dev["esn"] == "ESN-001"
    assert dev["_extras"]["esn"] == "ESN-001"
    assert dev["_extras"]["vrp_file"] == "SW.cc"
    assert dev["_extras"]["patch_file"] == "SW.pat"
