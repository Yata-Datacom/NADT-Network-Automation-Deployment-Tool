"""业务服务：校验、生成计划、Diff、批量操作、旧工作簿导入。"""

from __future__ import annotations

import csv
import difflib
import hashlib
import os
import re
import tempfile
from dataclasses import dataclass, field
from datetime import datetime

import nadt
from . import secrets
from .model import Project, Template

SEV_ORDER = {"error": 0, "warn": 1, "info": 2}


@dataclass
class Issue:
    severity: str
    code: str
    message: str
    device_index: int | None = None
    field: str = ""
    hint: str = ""

    @property
    def device_label(self) -> str:
        return "—" if self.device_index is None else f"第 {self.device_index + 1} 行"


@dataclass
class OutputFile:
    name: str
    content: str
    device_index: int | None = None
    status: str = "added"          # added / changed / unchanged / error
    old_content: str = ""
    error: str = ""
    is_lswnet: bool = False

    @property
    def changed(self) -> bool:
        return self.status in ("added", "changed")


@dataclass
class GenerationPlan:
    out_dir: str
    template_name: str
    files: list[OutputFile] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    generated_at: str = field(default_factory=lambda: datetime.now().strftime("%Y-%m-%d %H:%M:%S"))

    @property
    def added(self) -> int:
        return sum(1 for f in self.files if f.status == "added")

    @property
    def changed(self) -> int:
        return sum(1 for f in self.files if f.status == "changed")

    @property
    def unchanged(self) -> int:
        return sum(1 for f in self.files if f.status == "unchanged")

    @property
    def writable(self) -> int:
        return sum(1 for f in self.files if not f.error)


def active_secrets(project: Project) -> list[str]:
    return [c.value for c in project.credentials.values() if c.value]


def mask_secrets(text: str, project: Project) -> str:
    if not project.settings.mask_secrets:
        return text
    values = sorted(active_secrets(project), key=len, reverse=True)
    for value in values:
        if value and len(value) >= 4:
            text = text.replace(value, "••••••")
    return text


def diff_text(old: str, new: str, name: str, project: Project) -> str:
    old_m = mask_secrets(old or "", project)
    new_m = mask_secrets(new or "", project)
    old_lines = old_m.replace("\r\n", "\n").replace("\r", "\n").splitlines()
    new_lines = new_m.replace("\r\n", "\n").replace("\r", "\n").splitlines()
    diff = difflib.unified_diff(old_lines, new_lines, fromfile=f"现有 {name}",
                                tofile=f"新生成 {name}", lineterm="", n=3)
    return "\n".join(diff)


def diff_stats(old: str, new: str) -> tuple[int, int]:
    added = removed = 0
    old_lines = (old or "").replace("\r\n", "\n").splitlines()
    new_lines = (new or "").replace("\r\n", "\n").splitlines()
    for line in difflib.unified_diff(old_lines, new_lines, lineterm="", n=0):
        if line.startswith("+") and not line.startswith("+++"):
            added += 1
        elif line.startswith("-") and not line.startswith("---"):
            removed += 1
    return added, removed


class Validator:
    """项目级校验：重复、模板变量、文件名、凭据。"""

    def __init__(self, project: Project, template_name: str | None = None):
        self.project = project
        self.tpl = project.template(template_name)

    def validate(self) -> list[Issue]:
        p = self.project
        tpl = self.tpl
        issues: list[Issue] = []
        devices = p.target_devices()

        ext = (p.settings.extension or "cfg").lstrip(".")
        if not re.fullmatch(r"[A-Za-z0-9]{1,10}", ext):
            issues.append(Issue("error", "ext", f"后缀不合法: {ext}", hint="只用字母/数字，例如 cfg"))
        if not tpl.text.strip():
            issues.append(Issue("error", "template-empty", f"模板 {tpl.name} 为空"))
        if not devices:
            issues.append(Issue("warn", "no-devices", "当前范围没有设备，无法生成"))
            return self._sorted(issues)

        used_hosts: dict[str, int] = {}
        used_esn: dict[str, int] = {}
        used_names: dict[str, int] = {}
        rule = p.settings.filename_rule
        single = (rule == "__single__")
        vars_ = set(p.template_vars(tpl.text))
        secrets_by_key = p.credentials

        for i, dev in enumerate(devices):
            host = (dev.get("hostname") or "").strip()
            if not host:
                issues.append(Issue("error", "host-empty", "缺主机名", i, "hostname",
                                    "生成文件名与 lswnet 都需要主机名"))
            elif host in used_hosts:
                issues.append(Issue("error", "host-dup",
                                    f"主机名重复: {host}（与第 {used_hosts[host] + 1} 行冲突）",
                                    i, "hostname"))
            else:
                used_hosts[host] = i

            esn = (dev.get("esn") or "").strip()
            if esn:
                if esn in used_esn:
                    issues.append(Issue("warn", "esn-dup",
                                        f"ESN 重复: {esn}（与第 {used_esn[esn] + 1} 行冲突）",
                                        i, "esn"))
                else:
                    used_esn[esn] = i
            elif p.settings.generate_lswnet:
                issues.append(Issue("info", "esn-missing", "缺 ESN，不会写入 lswnet 映射表",
                                    i, "esn"))

            try:
                p.render(dev, tpl.name)
            except Exception as e:
                field = ""
                m = re.search(r"模板缺少变量[:：]\s*([^,\s，]+)", str(e))
                if m:
                    field = m.group(1)
                issues.append(Issue("error", "render", str(e), i, field,
                                    "检查模板占位符与表格参数名是否一致"))
                continue

            if not single:
                try:
                    name = p.filename_base(dev)
                    key = name.lower()
                    if key in used_names:
                        issues.append(Issue("error", "name-dup",
                                            f"文件名重复: {name}.{ext}（与第 {used_names[key] + 1} 行冲突）",
                                            i, rule))
                    else:
                        used_names[key] = i
                except Exception as e:
                    issues.append(Issue("error", "filename", str(e), i, rule))

        for key in sorted(vars_):
            if key in ("vlan_batch", "sysname"):
                continue
            if key in secrets_by_key and not secrets_by_key[key].value:
                issues.append(Issue("warn", "secret-empty",
                                    f"凭据 {key} 尚未填写（模板会渲染为空）", None, key))
            elif key in p.variables and not str(p.variables[key]).strip():
                issues.append(Issue("warn", "variable-empty",
                                    f"项目变量 {key} 为空（模板会渲染为空）", None, key))
        for key, cred in p.credentials.items():
            if cred.value and secrets.protection_kind(cred.to_dict().get("value", "")) == "plain":
                issues.append(Issue("info", "secret-plain",
                                    f"凭据 {key} 以明文保存（当前平台无 DPAPI）", None, key))
                break
        if not p.templates:
            issues.append(Issue("error", "no-template", "项目没有模板"))
        return self._sorted(issues)

    @staticmethod
    def _sorted(issues: list[Issue]) -> list[Issue]:
        return sorted(issues, key=lambda x: (SEV_ORDER.get(x.severity, 9),
                                             x.device_index if x.device_index is not None else -1,
                                             x.code))


class Generator:
    """预生成（不写盘）→ 查看 Diff → 确认后写盘。"""

    def __init__(self, project: Project, template_name: str | None = None,
                 out_dir: str | None = None):
        self.project = project
        self.tpl = project.template(template_name)
        self.out_dir = out_dir or project.resolve_output_dir(self.tpl.name)

    def plan(self) -> GenerationPlan:
        p = self.project
        tpl = self.tpl
        plan = GenerationPlan(out_dir=self.out_dir, template_name=tpl.name)
        targets = p.target_devices()
        if not targets:
            plan.errors.append("当前范围没有设备")
            return plan
        eol = "\r\n" if p.settings.eol == "crlf" else "\n"
        single = p.settings.filename_rule == "__single__"
        used: dict[str, int] = {}
        merged: list[tuple[str, str]] = []
        for i, dev in enumerate(targets):
            host = dev.get("hostname") or dev.get("esn") or f"设备{i + 1}"
            try:
                cfg = p.render(dev, tpl.name)
            except Exception as e:
                plan.errors.append(f"{host}: {e}")
                plan.files.append(OutputFile(name=f"{host}.ERROR", content="",
                                             device_index=i, status="error", error=str(e)))
                continue
            if single:
                merged.append((host, cfg))
                continue
            try:
                base = p.filename_base(dev)
            except Exception as e:
                plan.errors.append(f"{host}: {e}")
                plan.files.append(OutputFile(name=f"{host}.ERROR", content="",
                                             device_index=i, status="error", error=str(e)))
                continue
            fname = f"{base}.{p.settings.extension.lstrip('.')}"
            key = fname.lower()
            if key in used:
                msg = f"文件名重复: {fname}（与第 {used[key] + 1} 行冲突）"
                plan.errors.append(msg)
                plan.files.append(OutputFile(name=fname, content="", device_index=i,
                                             status="error", error=msg))
                continue
            used[key] = i
            content = self._with_eol(cfg, eol)
            old = self._read_existing(fname)
            status = self._status(old, content)
            plan.files.append(OutputFile(name=fname, content=content, old_content=old,
                                         device_index=i, status=status))

        if single and merged:
            safe_stem = re.sub(r'[\\/:*?"<>|]+', "_", tpl.stem)
            fname = f"{safe_stem}.{p.settings.extension.lstrip('.')}"
            body = "\n\n".join(c.rstrip("\n") for _h, c in merged) + "\n"
            content = self._with_eol(body, eol)
            old = self._read_existing(fname)
            plan.files.append(OutputFile(name=fname, content=content, old_content=old,
                                         device_index=None, status=self._status(old, content)))
            plan.warnings.append(f"单文件模式：{len(merged)} 台设备合并为 1 个文件")

        if p.settings.generate_lswnet and not single:
            ext = p.settings.extension.lstrip(".")
            if ext == "cfg":
                lswnet = nadt.build_lswnet(targets)
                if lswnet.strip():
                    content = self._with_eol(lswnet, eol)
                    old = self._read_existing(p.settings.bootfile)
                    plan.files.append(OutputFile(name=p.settings.bootfile, content=content,
                                                 old_content=old, status=self._status(old, content),
                                                 is_lswnet=True))
                else:
                    plan.warnings.append("没有可用 ESN，已跳过 lswnet.cfg")
            else:
                plan.warnings.append(f"lswnet.cfg 只适用于 .cfg，已跳过（当前后缀 {ext}）")
        return plan

    def _read_existing(self, name: str) -> str:
        path = os.path.join(self.out_dir, name)
        try:
            with open(path, "r", encoding="utf-8", errors="replace", newline="") as f:
                return f.read()
        except OSError:
            return ""

    @staticmethod
    def _status(old: str, new: str) -> str:
        if not old:
            return "added"
        return "unchanged" if old == new else "changed"

    @staticmethod
    def _with_eol(text: str, eol: str) -> str:
        return text.replace("\r\n", "\n").replace("\n", eol)

    def write(self, plan: GenerationPlan) -> dict:
        p = self.project
        os.makedirs(plan.out_dir, exist_ok=True)
        written = 0
        errors: list[str] = []
        manifest: dict[str, dict] = {}
        for f in plan.files:
            if f.error:
                continue
            manifest[f.name] = {
                "sha256": hashlib.sha256(f.content.encode("utf-8")).hexdigest(),
                "status": f.status,
                "device_index": f.device_index,
                "is_lswnet": f.is_lswnet,
            }
            if f.status == "unchanged":
                continue
            path = os.path.join(plan.out_dir, f.name)
            try:
                fd, tmp = tempfile.mkstemp(prefix=".nadt-out-", dir=plan.out_dir)
                with os.fdopen(fd, "w", encoding="utf-8", newline="") as fh:
                    fh.write(f.content)
                os.replace(tmp, path)
                written += 1
            except OSError as e:
                errors.append(f"{f.name}: {e}")
        p.last_generation = {
            "time": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "out_dir": plan.out_dir,
            "template": plan.template_name,
            "files": manifest,
        }
        p.mark_dirty()
        return {"written": written, "errors": errors, "manifest": manifest}


# ── 批量操作 ──
def fill_down(project: Project, indices: list[int], columns: list[str], value: str) -> int:
    n = 0
    for i in indices:
        if 0 <= i < len(project.devices):
            for c in columns:
                project.cell_set(project.devices[i], c, value)
                n += 1
    if n:
        project.mark_dirty()
    return n


def sequence(project: Project, indices: list[int], column: str, start: int, step: int,
             padding: int = 0, prefix: str = "", suffix: str = "") -> int:
    n = 0
    for k, i in enumerate(indices):
        if not (0 <= i < len(project.devices)):
            continue
        val = start + k * step
        s = f"{val:0{padding}d}" if padding else str(val)
        project.cell_set(project.devices[i], column, f"{prefix}{s}{suffix}")
        n += 1
    if n:
        project.mark_dirty()
    return n


def duplicate_rows(project: Project, indices: list[int]) -> list[int]:
    import copy
    wanted = set(indices)
    result = []
    new_indices = []
    for i, dev in enumerate(project.devices):
        result.append(dev)
        if i not in wanted:
            continue
        clone = copy.deepcopy(dev)
        host = (clone.get("hostname") or "device").strip()
        clone["hostname"] = f"{host}-copy"
        clone["esn"] = ""
        extras = clone.setdefault("_extras", {})
        extras["sysname"] = clone["hostname"]
        if "esn" in extras:
            extras["esn"] = ""
        clone["note"] = (clone.get("note") or "") + " (copy)"
        result.append(clone)
        new_indices.append(len(result) - 1)
    if new_indices:
        project.devices = result
        project.mark_dirty()
    return new_indices


def find_replace(project: Project, indices: list[int], columns: list[str],
                 find: str, replace: str, regex: bool = False,
                 case_sensitive: bool = False) -> int:
    if not find:
        return 0
    count = 0
    flags = 0 if case_sensitive else re.IGNORECASE
    pattern = re.compile(find if regex else re.escape(find), flags)
    for i in indices:
        if not (0 <= i < len(project.devices)):
            continue
        dev = project.devices[i]
        for c in columns:
            old = project.cell_get(dev, c)
            new = pattern.sub(replace, old)
            if new != old:
                project.cell_set(dev, c, new)
                count += 1
    if count:
        project.mark_dirty()
    return count


# ── 导入 / 导出 ──
def import_legacy_workbook(project: Project, path: str) -> dict:
    templates, devices, warnings = nadt.parse_xlsm_templates(path)
    added = 0
    for name, text in templates.items():
        fname = re.sub(r'[\\/:*?"<>|]+', "_", name) + ".cfg"
        if fname not in project.templates:
            project.templates[fname] = Template(name=fname, text=text, type_name=name)
            project.type_map[name.upper()] = fname
            added += 1
    if templates:
        first = re.sub(r'[\\/:*?"<>|]+', "_", next(iter(templates))) + ".cfg"
        project.settings.default_template = first
    project.devices.extend(devices)
    project.mark_dirty()
    return {"templates": added, "devices": len(devices), "warnings": warnings,
            "default": project.settings.default_template}


def import_workbook_as_templates(project: Project, path: str) -> dict:
    templates, devices, warnings = nadt.parse_xlsm_templates(path)
    for name, text in templates.items():
        fname = re.sub(r'[\\/:*?"<>|]+', "_", name) + ".cfg"
        if fname not in project.templates:
            project.templates[fname] = Template(name=fname, text=text, type_name=name)
            project.type_map[name.upper()] = fname
    project.mark_dirty()
    return {"templates": len(templates), "devices": len(devices), "warnings": warnings}


def read_table(path: str) -> tuple[list[str], list[list[str]]]:
    """读取 CSV/Excel 第一行表头 + 数据，全部转字符串。"""
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.active
        rows = [[("" if c is None else str(c)) for c in row]
                for row in ws.iter_rows(values_only=True)]
    else:
        with open(path, "r", encoding="utf-8-sig", errors="replace", newline="") as f:
            rows = [[str(c) for c in row] for row in csv.reader(f)]
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        return [], []
    width = max(len(r) for r in rows)
    header = [str(c).strip() for c in rows[0]] + [f"列{i}" for i in range(len(rows[0]), width)]
    data = [r + [""] * (width - len(r)) for r in rows[1:]]
    return header, data


def import_rows(project: Project, header: list[str], rows: list[list[str]],
                mapping: dict[int, str], template_name: str | None = None) -> int:
    """mapping: 源列索引 → 项目列名（hostname/esn/模板变量…）。"""
    count = 0
    for row in rows:
        dev = {k: "" for k in nadt.CSV_HEADER}
        dev["_extras"] = {}
        for idx, col in mapping.items():
            if col == "忽略" or idx >= len(row):
                continue
            project.cell_set(dev, col, row[idx])
        if not any(str(v).strip() for v in dev.values() if isinstance(v, str)) and \
                not any(dev.get("_extras", {}).values()):
            continue
        if not dev.get("type") and template_name:
            dev["type"] = project.template(template_name).stem
        project.devices.append(dev)
        count += 1
    if count:
        project.mark_dirty()
    return count


def export_project_csv(project: Project) -> str:
    return nadt.devices_to_csv(project.devices)


def export_workbook(project: Project, path: str):
    """把项目导出成旧 xlsm 兼容的 xlsx（每个模板一个 sheet + 设备行）。"""
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, tpl in project.templates.items():
        ws = wb.create_sheet(os.path.splitext(name)[0][:31])
        ws["A1"] = f"【{tpl.stem}】"
        ws["B1"] = "文本备注"
        cols = project.columns(name)
        for c, col in enumerate(cols, start=3):
            ws.cell(row=1, column=c, value="{{" + col + "}}")
        devs = [d for d in project.devices
                if not d.get("type") or d.get("type").lower() == tpl.stem.lower()]
        for r, dev in enumerate(devs, start=2):
            for c, col in enumerate(cols, start=3):
                ws.cell(row=r, column=c, value=project.cell_get(dev, col))
        for r, line in enumerate(tpl.text.splitlines(), start=2):
            ws.cell(row=r, column=1, value=line)
        ws.column_dimensions["A"].width = 60
    wb.save(path)
